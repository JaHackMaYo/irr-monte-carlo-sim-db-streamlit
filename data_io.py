import math
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

def load_monthly_profile_from_excel(path, profile_name):
    """Load an operating-month profile from the first worksheet of an Excel file.

    Accepted layouts are either one numeric column, or a first column containing
    dates/month indexes and a second numeric column. A header row is optional.
    Values are returned in worksheet order and are applied from COD onward.
    """
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError(f"{profile_name}: the workbook is empty.")
    values = []
    for row_number, row in enumerate(rows, 1):
        non_empty = [cell for cell in row if cell is not None and str(cell).strip() != ""]
        if not non_empty:
            continue
        candidates = non_empty[1:] if len(non_empty) >= 2 else non_empty
        numeric = None
        for cell in reversed(candidates):
            if isinstance(cell, bool):
                continue
            try:
                numeric = float(str(cell).strip().replace(",", "."))
                break
            except (TypeError, ValueError):
                continue
        if numeric is None:
            if not values:
                continue  # optional header row
            raise ValueError(f"{profile_name}: no numeric value found at row {row_number}.")
        if not math.isfinite(numeric):
            raise ValueError(f"{profile_name}: non-finite value at row {row_number}.")
        values.append(numeric)
    if not values:
        raise ValueError(f"{profile_name}: no numeric monthly values were found.")
    return values

def export_generated_flows_to_excel(input_parameters, generated_flows, path, input_rows):
    """Export GUI inputs and one generated iteration to a simple workbook."""
    (
        irr,
        actual_eoh,
        pbp,
        capex_flows,
        eoh,
        unavailabilities,
        energy_prices,
        revenues_flows,
        taxes_flows,
        opexes,
    ) = generated_flows
    series = {
        "CAPEX (MEUR)": list(capex_flows),
        "EOH (h)": list(eoh),
        "Unavailability": list(unavailabilities),
        "Energy Price (EUR/MWh)": list(energy_prices),
        "Gross Revenues (MEUR)": list(revenues_flows),
        "OPEX (MEUR)": list(opexes),
        "Taxes (MEUR)": list(taxes_flows),
    }
    lengths = {len(values) for values in series.values()}
    if len(lengths) != 1:
        raise ValueError("Generated monthly flow series have different lengths.")
    total_months = lengths.pop()
    operational_life = int(input_parameters["operational_life"])
    actual_time_to_cod = total_months - operational_life * 12
    if actual_time_to_cod < 0:
        raise ValueError("Unable to derive a valid Time to COD from generated flows.")
    net_cash_flows = []
    cumulative_cash_flows = []
    cumulative = 0.0
    for capex, revenue, opex, tax in zip(capex_flows, revenues_flows, opexes, taxes_flows):
        net_cash_flow = -float(capex) + float(revenue) - float(opex) - float(tax)
        cumulative += net_cash_flow
        net_cash_flows.append(net_cash_flow)
        cumulative_cash_flows.append(cumulative)
    series["Net Cash Flow (MEUR)"] = net_cash_flows
    series["Cumulative Cash Flow (MEUR)"] = cumulative_cash_flows
    workbook = Workbook()
    inputs_sheet = workbook.active
    inputs_sheet.title = "Inputs"
    flows_sheet = workbook.create_sheet("Monthly Flows")
    inputs_sheet.append(["Input", "Value"])
    for label, value in input_rows:
        inputs_sheet.append([label, value])
    inputs_sheet.append(["Actual CAPEX (M€)", float(sum(capex_flows))])
    inputs_sheet.append(["Actual Time to COD (m)", actual_time_to_cod])
    inputs_sheet.append(["IRR (%)", float(irr)])
    inputs_sheet.append(["Actual EOH (h)", float(actual_eoh)])
    inputs_sheet.append(["PBP (y)", float(pbp)])
    inputs_sheet.column_dimensions["A"].width = 38
    inputs_sheet.column_dimensions["B"].width = 24
    inputs_sheet.freeze_panes = "A2"
    inputs_sheet.auto_filter.ref = f"A1:B{inputs_sheet.max_row}"
    start_year = int(input_parameters.get("start_year", 2026))
    start_month = int(input_parameters.get("start_month", 1))
    if not 1 <= start_month <= 12:
        raise ValueError("Start Month must be between 1 and 12.")
    flow_dates = [
        date(
            start_year + (start_month - 1 + month) // 12,
            (start_month - 1 + month) % 12 + 1,
            15,
        )
        for month in range(total_months)
    ]
    headers = ["Month", "Date", "Phase", "Operating Month"] + list(series)
    flows_sheet.append(headers)
    for month in range(total_months):
        construction = month < actual_time_to_cod
        row = [
            month + 1,
            flow_dates[month],
            "Construction" if construction else "Operation",
            0 if construction else month - actual_time_to_cod + 1,
        ]
        row.extend(float(values[month]) for values in series.values())
        flows_sheet.append(row)
    flows_sheet.freeze_panes = "A2"
    flows_sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{flows_sheet.max_row}"
    for column in range(1, len(headers) + 1):
        header = headers[column - 1]
        flows_sheet.column_dimensions[get_column_letter(column)].width = max(
            12, min(28, len(header) + 3)
        )
    money_format = "#,##0.000;[Red](#,##0.000);-"
    number_format = "#,##0.00;[Red](#,##0.00);-"
    for row in range(2, flows_sheet.max_row + 1):
        flows_sheet.cell(row, 2).number_format = "dd/mm/yyyy"
        flows_sheet.cell(row, 5).number_format = money_format
        flows_sheet.cell(row, 6).number_format = number_format
        flows_sheet.cell(row, 7).number_format = "0.00%"
        flows_sheet.cell(row, 8).number_format = "€ #,##0.00;[Red](€ #,##0.00);-"
        for column in range(9, len(headers) + 1):
            flows_sheet.cell(row, column).number_format = money_format
    workbook.save(path)