# python -m streamlit run streamlit_app_db_.py


import sys

if sys.platform == "win32":
    try:
        import truststore

        truststore.inject_into_ssl()
    except ImportError:
        pass

import inspect
import io
import math
from datetime import date

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from scipy.stats import linregress
import plotly.graph_objects as go
from plotly.subplots import make_subplots

from config import *
from core import (
    capex_generator,
    distribute_capex,
    energy_price_generator,
    generate_flows,
    my_irr_mc_function,
    opex_generator,
    time_to_cod_generator,
    unavailability_generator,
)
from models import MonteCarloResults
from database import (
    delete_preset,
    get_preset,
    list_presets,
    save_preset,
    test_connection,
    update_preset,
)

st.set_page_config(page_title="IRR Monte Carlo DB", page_icon="📊", layout="wide", initial_sidebar_state="expanded")

BG = "#F4F6F4"
CARD = "#FFFFFF"
HEADER = "#26352D"
PRIMARY = "#486B57"
ACCENT = "#6F8F78"
MUTED = "#6B736E"
TEXT = "#202622"
GREEN = "#2F6048"
LIGHT_GREEN = "#E4EEE7"
BORDER = "#C9D2CC"
GRID = "#DDE3DF"

# Result and consultation charts use a separate, internally consistent palette.
CHART_BLUE = "#3E6C8A"
CHART_BLUE_LIGHT = "#8AB6CC"
CHART_ORANGE = "#D17A45"
CHART_RED = "#B44C4C"
CHART_PURPLE = "#77659A"
CHART_GOLD = "#C69A3D"
CHART_TEAL = "#3F837A"
CHART_SLATE = "#596A76"
BAR_OUTLINE = "#263238"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

PLOTLY_CONFIG = {
    "displaylogo": False,
    "scrollZoom": True,
    "responsive": True,
    "toImageButtonOptions": {"format": "png", "filename": "irr_chart", "scale": 2},
}


def style_figure(fig, title, x_title=None, y_title=None, height=430, hovermode="x unified"):
    fig.update_layout(
        title=dict(text=title, x=0.01, xanchor="left"),
        height=height,
        margin=dict(l=45, r=30, t=75, b=45),
        paper_bgcolor=CARD,
        plot_bgcolor=CARD,
        font=dict(family="Segoe UI", color=TEXT),
        hovermode=hovermode,
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
    )
    fig.update_xaxes(title=x_title, showgrid=True, gridcolor=GRID, zeroline=False)
    fig.update_yaxes(title=y_title, showgrid=True, gridcolor=GRID, zeroline=False)
    return fig


def show_plotly(fig, key=None):
    st.plotly_chart(fig, width="stretch", config=PLOTLY_CONFIG, key=key)


def interactive_histogram(values, reference, title, xlabel, bins=35, discrete=False):
    values = finite(values)
    if values.size == 0:
        return go.Figure()
    if discrete:
        values = np.rint(values).astype(int)
        lo, hi = int(values.min()), int(values.max())
        xbins = dict(start=lo - 0.5, end=hi + 0.5, size=1)
    else:
        width = (values.max() - values.min()) / max(int(bins), 1)
        xbins = dict(size=width) if width > 0 else None
    fig = go.Figure()
    fig.add_trace(go.Histogram(
        x=values, name="Frequency", marker=dict(color=CHART_BLUE, line=dict(color=BAR_OUTLINE, width=1.1)), opacity=0.84,
        xbins=xbins, hovertemplate="Range: %{x}<br>Frequency: %{y}<extra></extra>"
    ))
    stats = [
        (reference, "Reference", CHART_SLATE, "solid"),
        (float(np.mean(values)), "Mean", CHART_ORANGE, "solid"),
        (float(np.median(values)), "Median", CHART_TEAL, "dash"),
        (float(np.percentile(values, 5)), "P5", CHART_PURPLE, "dot"),
        (float(np.percentile(values, 95)), "P95", CHART_RED, "dot"),
    ]
    for value, label, color, dash in stats:
        fig.add_vline(x=value, line_width=2, line_color=color, line_dash=dash)
        fig.add_trace(go.Scatter(x=[None], y=[None], mode="lines", name=f"{label}: {value:.2f}", line=dict(color=color, dash=dash)))
    return style_figure(fig, title, xlabel, "Frequency", 470, "x")


DISTRIBUTIONS = {
    "EOH": ("iterations_eoh", "EOH", False),
    "EOH Variation (%)": ("eoh_variations", "EOH Variation (%)", False),
    "Time to COD (m)": ("iterations_time_to_cod", "Time to COD (m)", True),
    "Delay (m)": ("delays", "Delay (m)", True),
    "CAPEX (M€)": ("iterations_capex", "CAPEX (M€)", False),
    "CAPEX Variation (%)": ("capex_variations", "CAPEX Variation (%)", False),
    "IRR Variation (bps)": ("irr_variations", "IRR Variation (bps)", False),
}
SCATTER_VARIABLES = {
    "IRR (%)": ("iterations_irr", "IRR (%)", 100.0),
    "IRR Variation (bps)": ("irr_variations", "IRR Variation (bps)", 1.0),
    "EOH": ("iterations_eoh", "EOH", 1.0),
    "EOH Variation (%)": ("eoh_variations", "EOH Variation (%)", 1.0),
    "Time to COD (m)": ("iterations_time_to_cod", "Time to COD (m)", 1.0),
    "Delay (m)": ("delays", "Delay (m)", 1.0),
    "CAPEX (M€)": ("iterations_capex", "CAPEX (M€)", 1.0),
    "CAPEX Variation (%)": ("capex_variations", "CAPEX Variation (%)", 1.0),
    "PBP (y)": ("iterations_bpb", "PBP (y)", 1.0),
}

st.markdown(f"""
<style>
.stApp {{background:{BG}; color:{TEXT};}}
[data-testid="stSidebar"] {{background:{CARD}; border-right:1px solid {BORDER};}}
.block-container {{padding-top:1.2rem; padding-bottom:2rem; max-width:1900px;}}
.hero {{background:{HEADER}; color:white; padding:18px 24px; border-radius:10px; margin-bottom:14px;}}
.hero h1 {{margin:0; font-size:2rem;}}
.hero p {{margin:4px 0 0 0; color:#E9EEEA;}}
.status-ready {{background:{LIGHT_GREEN}; color:{GREEN}; border:1px solid #B7CBBB; padding:7px 11px; border-radius:8px; font-weight:600;}}
.status-error {{background:#F1F2F1; color:#4B514D; border:1px solid {BORDER}; padding:7px 11px; border-radius:8px; font-weight:600;}}
div[data-testid="stMetric"] {{background:{CARD}; border:1px solid {BORDER}; padding:13px; border-radius:9px;}}
div[data-testid="stMetricValue"] {{color:{GREEN};}}
.stButton>button, .stDownloadButton>button {{background:{CARD}; color:{GREEN}; border:1px solid {GREEN}; border-radius:8px;}}
.stButton>button:hover, .stDownloadButton>button:hover {{background:{LIGHT_GREEN}; color:{HEADER}; border-color:{HEADER};}}
.stButton>button[kind="primary"] {{background:{GREEN}; color:white; border-color:{GREEN};}}
.stButton>button[kind="primary"]:hover {{background:{HEADER}; color:white; border-color:{HEADER};}}
button[data-baseweb="tab"] {{background:#ECEFED; color:{MUTED}; border-radius:8px 8px 0 0;}}
button[data-baseweb="tab"][aria-selected="true"] {{background:{LIGHT_GREEN}; color:{GREEN}; font-weight:700;}}
[data-baseweb="slider"] [role="slider"] {{background:{GREEN}; border-color:{GREEN};}}
[data-baseweb="slider"] > div > div {{background:{ACCENT};}}
[data-baseweb="checkbox"] div[aria-checked="true"] {{background:{GREEN};}}
[data-testid="stExpander"] {{background:{CARD}; border:1px solid {BORDER}; border-radius:8px;}}
[data-baseweb="select"] > div, [data-baseweb="input"] > div {{background:{CARD}; border-color:{BORDER};}}
</style>
""", unsafe_allow_html=True)


def init_state():
    defaults = {
        "results": None,
        "single_iteration": None,
        "last_params": None,
        "current_params": None,
        "monte_carlo_params": None,
        "single_iteration_params": None,
        "status": "Ready",
        "error": None,
        "energy_price_profile": None,
        "unavailability_profile": None,
        "opex_profile": None,
        "mean_eoh": list(MEAN_EOH),
        "std_eoh": list(STD_EOH),
        "selected_preset_id": None,
        "selected_preset_name": "",
        "selected_preset_description": "",
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def finite(values):
    values = np.asarray(values, dtype=float)
    return values[np.isfinite(values)]


def parse_profile(uploaded, profile_name):
    if uploaded is None:
        return None
    wb = load_workbook(io.BytesIO(uploaded.getvalue()), data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        values = []
        for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = [v for v in row if v is not None and str(v).strip()]
            if not cells:
                continue
            candidates = cells[1:] if len(cells) >= 2 else cells
            numeric = None
            for cell in reversed(candidates):
                try:
                    numeric = float(str(cell).strip().replace(",", "."))
                    break
                except (TypeError, ValueError):
                    pass
            if numeric is None:
                if not values:
                    continue
                raise ValueError(f"{profile_name}: no numeric value at row {row_number}")
            if not math.isfinite(numeric):
                raise ValueError(f"{profile_name}: non-finite value at row {row_number}")
            values.append(numeric)
    finally:
        wb.close()
    if not values:
        raise ValueError(f"{profile_name}: no numeric values found")
    return values


def profile_workbook(label, unit, values):
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Profile"
    ws.append(["Operating Month", f"{label} ({unit})"])
    for i, value in enumerate(values, 1):
        ws.append([i, float(value)])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 24
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


BOOLEAN_INPUT_KEYS = {
    "stochastic_capex",
    "stochastic_delay",
    "stochastic_eoh",
}
INTEGER_INPUT_KEYS = {
    "iterations",
    "start_year",
    "start_month",
    "nominal_time_to_cod",
    "operational_life",
    "capex_depreciation_time",
    "delay_min",
    "delay_max",
    "irr_max_iter",
}
UNAVAILABILITY_INPUT_KEYS = (
    "initial_unavailability",
    "steady_state_unavailability",
    "ramp_down_months",
    "aging_degradation",
)


def _parse_boolean(value, label):
    """Parse a real Excel boolean without treating arbitrary text as True."""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {"true", "yes", "y", "1", "si", "sì"}:
            return True
        if normalized in {"false", "no", "n", "0"}:
            return False
    raise ValueError(f"{label} must be TRUE or FALSE.")


def _parse_number(value, label, integer=False):
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{label} must be numeric.")
    if isinstance(value, str):
        value = value.strip().replace(",", ".")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be numeric.") from exc
    if not math.isfinite(number):
        raise ValueError(f"{label} must be finite.")
    if integer:
        rounded = round(number)
        if abs(number - rounded) > 1e-9:
            raise ValueError(f"{label} must be an integer.")
        return int(rounded)
    return number


def _read_profile_column(sheet, column, label):
    values = []
    for row_number in range(5, sheet.max_row + 1):
        month = sheet.cell(row_number, 1).value
        value = sheet.cell(row_number, column).value
        if month is None and value is None:
            continue
        if month is None:
            raise ValueError(f"{label}: missing operating month at row {row_number}.")
        month_number = _parse_number(month, f"Operating month at row {row_number}", True)
        if month_number != len(values) + 1:
            raise ValueError(
                f"{label}: operating months must be consecutive from 1."
            )
        if value is None:
            raise ValueError(f"{label}: missing value at row {row_number}.")
        values.append(_parse_number(value, f"{label} at row {row_number}"))
    return values


def params_workbook(params):
    """Export scalar inputs, explicit model settings, EOH data and profiles."""
    workbook = Workbook()
    inputs_sheet = workbook.active
    inputs_sheet.title = "Inputs"
    inputs_sheet.append(["Input", "Value"])

    excluded = {
        "energy_price_profile",
        "opex_profile",
        "unavailability_profile",
        "mean_eoh",
        "std_eoh",
        "unavailability_parameters",
    }
    for key, value in params.items():
        if key in excluded:
            continue
        inputs_sheet.append([key, value])

    unavailability = list(params["unavailability_parameters"])
    if len(unavailability) != 4:
        raise ValueError("Unavailability parameters must contain four values.")
    for key, value in zip(UNAVAILABILITY_INPUT_KEYS, unavailability):
        inputs_sheet.append([key, value])

    inputs_sheet.append([])
    inputs_sheet.append(["EOH Month", "Mean EOH", "EOH CV"])
    mean_eoh = list(params["mean_eoh"])
    std_eoh = list(params["std_eoh"])
    if len(mean_eoh) != 12 or len(std_eoh) != 12:
        raise ValueError("Mean EOH and EOH CV must contain exactly 12 values.")
    for month, mean_value, cv_value in zip(MONTHS, mean_eoh, std_eoh):
        inputs_sheet.append([month, float(mean_value), float(cv_value)])

    inputs_sheet.freeze_panes = "A2"
    inputs_sheet.column_dimensions["A"].width = 38
    inputs_sheet.column_dimensions["B"].width = 22
    inputs_sheet.column_dimensions["C"].width = 18

    use_flags = {
        "energy_price": params.get("energy_price_profile") is not None,
        "opex": params.get("opex_profile") is not None,
        "unavailability": params.get("unavailability_profile") is not None,
    }
    life = int(params["operational_life"])
    energy_price_values = energy_price_generator(
        0,
        life,
        params["energy_price_at_cod"],
        params["annual_price_variation"],
        series=params.get("energy_price_profile"),
    )
    opex_values = opex_generator(
        0,
        life,
        params["opex"],
        params["capacity"],
        series=params.get("opex_profile"),
    )
    unavailability_values = np.asarray(
        unavailability_generator(
            0,
            life,
            *unavailability,
            series=params.get("unavailability_profile"),
        ),
        dtype=float,
    ) * 100.0

    profiles_sheet = workbook.create_sheet("Operating Profiles")
    profiles_sheet.append(
        [
            "Use Energy Price Series",
            "Use OPEX Series",
            "Use Unavailability Series",
        ]
    )
    profiles_sheet.append(
        [
            use_flags["energy_price"],
            use_flags["opex"],
            use_flags["unavailability"],
        ]
    )
    profiles_sheet.append([])
    profiles_sheet.append(
        [
            "Operating Month",
            "Energy Price (EUR/MWh)",
            "OPEX (M€/month)",
            "Unavailability (%)",
        ]
    )
    expected_months = life * 12
    for index in range(expected_months):
        profiles_sheet.append(
            [
                index + 1,
                float(energy_price_values[index]),
                float(opex_values[index]),
                float(unavailability_values[index]),
            ]
        )
    profiles_sheet.freeze_panes = "A5"
    profiles_sheet.column_dimensions["A"].width = 19
    profiles_sheet.column_dimensions["B"].width = 27
    profiles_sheet.column_dimensions["C"].width = 22
    profiles_sheet.column_dimensions["D"].width = 25

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def load_params_workbook(uploaded):
    """Load and validate the complete workbook exported by this application."""
    workbook = load_workbook(io.BytesIO(uploaded.getvalue()), data_only=True)
    try:
        required_sheets = {"Inputs", "Operating Profiles"}
        missing = required_sheets.difference(workbook.sheetnames)
        if missing:
            raise ValueError(f"Missing worksheet: {', '.join(sorted(missing))}.")

        inputs_sheet = workbook["Inputs"]
        imported_inputs = {}
        eoh_header_row = None
        for row_number in range(2, inputs_sheet.max_row + 1):
            key = inputs_sheet.cell(row_number, 1).value
            value = inputs_sheet.cell(row_number, 2).value
            if key == "EOH Month":
                eoh_header_row = row_number
                break
            if key is None:
                continue
            key = str(key).strip()
            if key in BOOLEAN_INPUT_KEYS:
                imported_inputs[key] = _parse_boolean(value, key)
            elif key in INTEGER_INPUT_KEYS:
                imported_inputs[key] = _parse_number(value, key, True)
            elif key in UNAVAILABILITY_INPUT_KEYS:
                imported_inputs[key] = _parse_number(
                    value,
                    key,
                    integer=(key == "ramp_down_months"),
                )
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                imported_inputs[key] = _parse_number(value, key)
            else:
                imported_inputs[key] = value

        if eoh_header_row is None:
            raise ValueError("The Inputs sheet does not contain the EOH table.")
        mean_eoh = []
        std_eoh = []
        for offset, expected_month in enumerate(MONTHS, 1):
            row_number = eoh_header_row + offset
            month = inputs_sheet.cell(row_number, 1).value
            if month != expected_month:
                raise ValueError(
                    f"EOH row {row_number}: expected month {expected_month}."
                )
            mean_eoh.append(
                _parse_number(
                    inputs_sheet.cell(row_number, 2).value,
                    f"Mean EOH for {expected_month}",
                )
            )
            std_eoh.append(
                _parse_number(
                    inputs_sheet.cell(row_number, 3).value,
                    f"EOH CV for {expected_month}",
                )
            )
        imported_inputs["mean_eoh"] = mean_eoh
        imported_inputs["std_eoh"] = std_eoh

        missing_unavailability = [
            key for key in UNAVAILABILITY_INPUT_KEYS if key not in imported_inputs
        ]
        if missing_unavailability:
            raise ValueError(
                "Missing unavailability settings: "
                + ", ".join(missing_unavailability)
            )
        imported_inputs["unavailability_parameters"] = [
            imported_inputs.pop(key) for key in UNAVAILABILITY_INPUT_KEYS
        ]

        profiles_sheet = workbook["Operating Profiles"]
        flags = {
            "energy_price": _parse_boolean(
                profiles_sheet.cell(2, 1).value, "Use Energy Price Series"
            ),
            "opex": _parse_boolean(
                profiles_sheet.cell(2, 2).value, "Use OPEX Series"
            ),
            "unavailability": _parse_boolean(
                profiles_sheet.cell(2, 3).value,
                "Use Unavailability Series",
            ),
        }
        profiles = {
            "energy_price": _read_profile_column(
                profiles_sheet, 2, "Energy Price series"
            ),
            "opex": _read_profile_column(profiles_sheet, 3, "OPEX series"),
            "unavailability": _read_profile_column(
                profiles_sheet, 4, "Unavailability series"
            ),
        }

        if any(value < 0 for value in profiles["energy_price"]):
            raise ValueError("Energy Price series cannot contain negative values.")
        if any(value < 0 for value in profiles["opex"]):
            raise ValueError("OPEX series cannot contain negative values.")
        if any(value < 0 or value > 100 for value in profiles["unavailability"]):
            raise ValueError(
                "Unavailability series values must be between 0 and 100%."
            )

        expected_months = int(imported_inputs["operational_life"]) * 12
        for key, values in profiles.items():
            if len(values) != expected_months:
                raise ValueError(
                    f"{key.replace('_', ' ').title()} series contains "
                    f"{len(values)} values; {expected_months} are required."
                )
        return imported_inputs, profiles, flags
    finally:
        workbook.close()


def input_defaults():
    """Return a complete valid parameter set for the downloadable template."""
    return {
        "project_name": PROJECT_NAME,
        "iterations": int(ITERATIONS),
        "capex_delay_correlation": float(CAPEX_DELAY_CORRELATION),
        "stochastic_capex": bool(STOCHASTIC_CAPEX),
        "stochastic_delay": bool(STOCHASTIC_DELAY),
        "stochastic_eoh": bool(STOCHASTIC_EOH),
        "nominal_capex": float(NOMINAL_CAPEX),
        "start_year": int(START_YEAR),
        "start_month": int(START_MONTH),
        "nominal_time_to_cod": int(NOMINAL_TIME_TO_COD),
        "operational_life": int(OPERATIONAL_LIFE),
        "capacity": float(CAPACITY),
        "mean_eoh": list(MEAN_EOH),
        "std_eoh": list(STD_EOH),
        "mean_eoh_variation": float(MEAN_EOH_VARIATION),
        "unavailability_parameters": list(UNAVAILABILITY_PARAMETERS),
        "opex": float(OPEX),
        "energy_price_at_cod": float(ENERGY_PRICE_AT_COD),
        "annual_price_variation": float(ANNUAL_PRICE_VARIATION),
        "energy_price_profile": None,
        "unavailability_profile": None,
        "opex_profile": None,
        "tax_rate": float(TAX_RATE),
        "capex_depreciation_ratio": float(CAPEX_DEPRECIATION_RATIO),
        "capex_depreciation_time": int(CAPEX_DEPRECIATION_TIME),
        "delay_min": int(DELAY_MIN),
        "delay_ml": float(DELAY_ML),
        "delay_max": int(DELAY_MAX),
        "delay_distribution_sharpness": float(DELAY_DISTRIBUTION_SHARPNESS),
        "capex_variation_min": float(CAPEX_VARIATION_MIN),
        "capex_variation_ml": float(CAPEX_VARIATION_ML),
        "capex_variation_max": float(CAPEX_VARIATION_MAX),
        "capex_distribution_sharpness": float(CAPEX_DISTRIBUTION_SHARPNESS),
        "capex_curve_inflection": float(CAPEX_CURVE_INFLECTION),
        "capex_curve_steepness": float(CAPEX_CURVE_STEEPNESS),
        "capex_stress": float(CAPEX_STRESS),
        "delay_stress": float(DELAY_STRESS),
        "eoh_stress": float(EOH_STRESS),
        "irr_guess": float(IRR_GUESS),
        "irr_tol": float(IRR_TOL),
        "irr_max_iter": int(IRR_MAX_ITER),
    }

def single_iteration_workbook(generated, params):
    """Export inputs, active profiles, visible results and all monthly series."""
    (
        irr,
        actual_eoh,
        pbp,
        capex,
        eoh,
        unavailability,
        price,
        revenues,
        taxes,
        opex,
    ) = generated

    capex = np.asarray(capex, dtype=float)
    eoh = np.asarray(eoh, dtype=float)
    unavailability = np.asarray(unavailability, dtype=float)
    price = np.asarray(price, dtype=float)
    revenues = np.asarray(revenues, dtype=float)
    taxes = np.asarray(taxes, dtype=float)
    opex = np.asarray(opex, dtype=float)

    lengths = {
        len(capex),
        len(eoh),
        len(unavailability),
        len(price),
        len(revenues),
        len(taxes),
        len(opex),
    }
    if len(lengths) != 1:
        raise ValueError("Single-iteration series have different lengths.")

    total_months = lengths.pop()
    operational_life = int(params["operational_life"])
    actual_time_to_cod = total_months - operational_life * 12
    if actual_time_to_cod < 0:
        raise ValueError("Unable to derive a valid actual Time to COD.")

    net_cash_flow = -capex + revenues - opex - taxes
    cumulative_cash_flow = np.cumsum(net_cash_flow)

    # Start from the same complete workbook used by the input template. This
    # preserves the exact input format, operating profiles and activation flags.
    workbook = load_workbook(io.BytesIO(params_workbook(params)))

    results_sheet = workbook.create_sheet("Iteration Results")
    results_sheet.append(["Result", "Value", "Unit"])
    results_sheet.append(["IRR", float(irr), "%"])
    results_sheet.append(["PBP", float(pbp), "years"])
    results_sheet.append(["Actual CAPEX", float(capex.sum()), "M€"])
    results_sheet.append(["Actual Time to COD", actual_time_to_cod, "months"])
    results_sheet.append(["Total EOH", float(actual_eoh), "hours"])
    results_sheet.freeze_panes = "A2"
    results_sheet.column_dimensions["A"].width = 28
    results_sheet.column_dimensions["B"].width = 18
    results_sheet.column_dimensions["C"].width = 14
    results_sheet["B2"].number_format = "0.00%"
    results_sheet["B3"].number_format = "0.00"
    results_sheet["B4"].number_format = "0.000"
    results_sheet["B5"].number_format = "0"
    results_sheet["B6"].number_format = "#,##0.00"

    flows_sheet = workbook.create_sheet("Monthly Flows")
    headers = [
        "Month",
        "Date",
        "Phase",
        "Operating Month",
        "CAPEX (M€)",
        "EOH (h)",
        "Unavailability",
        "Energy Price (EUR/MWh)",
        "Gross Revenues (M€)",
        "OPEX (M€)",
        "Taxes (M€)",
        "Net Cash Flow (M€)",
        "Cumulative Cash Flow (M€)",
    ]
    flows_sheet.append(headers)

    start_year = int(params["start_year"])
    start_month = int(params["start_month"])
    for index in range(total_months):
        flow_date = date(
            start_year + (start_month - 1 + index) // 12,
            (start_month - 1 + index) % 12 + 1,
            15,
        )
        construction = index < actual_time_to_cod
        flows_sheet.append(
            [
                index + 1,
                flow_date,
                "Construction" if construction else "Operation",
                0 if construction else index - actual_time_to_cod + 1,
                float(capex[index]),
                float(eoh[index]),
                float(unavailability[index]),
                float(price[index]),
                float(revenues[index]),
                float(opex[index]),
                float(taxes[index]),
                float(net_cash_flow[index]),
                float(cumulative_cash_flow[index]),
            ]
        )

    flows_sheet.freeze_panes = "A2"
    flows_sheet.auto_filter.ref = flows_sheet.dimensions
    widths = [10, 13, 15, 18, 16, 13, 18, 25, 23, 15, 15, 24, 30]
    for column, width in enumerate(widths, 1):
        flows_sheet.column_dimensions[
            flows_sheet.cell(1, column).column_letter
        ].width = width
    for row_number in range(2, flows_sheet.max_row + 1):
        flows_sheet.cell(row_number, 2).number_format = "dd/mm/yyyy"
        flows_sheet.cell(row_number, 5).number_format = "#,##0.000"
        flows_sheet.cell(row_number, 6).number_format = "#,##0.00"
        flows_sheet.cell(row_number, 7).number_format = "0.00%"
        flows_sheet.cell(row_number, 8).number_format = "#,##0.00"
        for column in range(9, 14):
            flows_sheet.cell(row_number, column).number_format = "#,##0.000"

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def validate(params):
    if params["iterations"] <= 0:
        raise ValueError("Iterations must be greater than zero")
    if params["nominal_capex"] <= 0 or params["nominal_time_to_cod"] <= 0 or params["operational_life"] <= 0:
        raise ValueError("CAPEX, Time to COD and Operational Life must be greater than zero")
    if not -1 <= params["capex_delay_correlation"] <= 1:
        raise ValueError("CAPEX-delay correlation must be between -1 and 1")
    if not params["delay_min"] <= params["delay_ml"] <= params["delay_max"]:
        raise ValueError("Minimum Delay <= Most Likely Delay <= Maximum Delay must hold")
    if not params["capex_variation_min"] <= params["capex_variation_ml"] <= params["capex_variation_max"]:
        raise ValueError("Minimum CAPEX variation <= Most Likely <= Maximum must hold")
    if any(v < 0 for v in params["std_eoh"]):
        raise ValueError("EOH coefficients of variation cannot be negative")
    if any(v < 0 for v in params["unavailability_parameters"]):
        raise ValueError("Unavailability parameters cannot be negative")
    return params


def input_panel(defaults=None):
    defaults = defaults or {}

    def d(name, fallback):
        return defaults.get(name, fallback)

    st.subheader("Inputs")
    general, operation, uncertainty, advanced = st.tabs(["General", "Operating models", "CAPEX, Time to COD and EOH", "Advanced"])
    with general:
        project_name = st.text_input("Project Name", d("project_name", PROJECT_NAME))
        capacity = st.number_input(
            "Capacity (MW)", min_value=0.0, value=float(d("capacity", CAPACITY))
        )
        iterations = st.number_input(
            "Iterations", min_value=1, value=int(d("iterations", ITERATIONS)), step=100
        )

        st.divider()
        stochastic_capex = st.toggle(
            "Simulate Stochastic CAPEX", bool(d("stochastic_capex", STOCHASTIC_CAPEX))
        )
        stochastic_delay = st.toggle(
            "Simulate Stochastic Delay", bool(d("stochastic_delay", STOCHASTIC_DELAY))
        )
        stochastic_eoh = st.toggle(
            "Simulate Stochastic EOH", bool(d("stochastic_eoh", STOCHASTIC_EOH))
        )

        st.divider()
        nominal_capex = st.number_input(
            "Nominal CAPEX (M€)", min_value=0.0, value=float(d("nominal_capex", NOMINAL_CAPEX))
        )

        st.divider()
        start_year = st.number_input(
            "Start Year", min_value=1, value=int(d("start_year", START_YEAR)), step=1
        )
        start_month = st.number_input(
            "Start Month (1 = Jan)",
            min_value=1,
            max_value=12,
            value=int(d("start_month", START_MONTH)),
            step=1,
        )
        nominal_time_to_cod = st.number_input(
            "Nominal Time to COD (m)",
            min_value=1,
            value=int(d("nominal_time_to_cod", NOMINAL_TIME_TO_COD)),
            step=1,
        )
        operational_life = st.number_input(
            "Operational Life (y)",
            min_value=1,
            value=int(d("operational_life", OPERATIONAL_LIFE)),
            step=1,
        )

        st.divider()
        tax_rate = st.number_input(
            "Tax Rate (1 = 100%)",
            min_value=0.0,
            value=float(d("tax_rate", TAX_RATE)),
            format="%.4f",
        )
        depreciation_ratio = st.number_input(
            "Depreciable CAPEX Ratio",
            min_value=0.0,
            max_value=1.0,
            value=float(d("capex_depreciation_ratio", CAPEX_DEPRECIATION_RATIO)),
            format="%.4f",
        )
        depreciation_time = st.number_input(
            "Depreciation Period (y)",
            min_value=1,
            value=int(d("capex_depreciation_time", CAPEX_DEPRECIATION_TIME)),
            step=1,
        )

        st.divider()
        st.markdown("#### Stress")
        capex_stress = st.number_input(
            "CAPEX Stress (%)", value=float(d("capex_stress", CAPEX_STRESS))
        )
        delay_stress = st.number_input(
            "Delay Stress (m)", value=float(d("delay_stress", DELAY_STRESS))
        )
        eoh_stress = st.number_input(
            "EOH Stress (%)", value=float(d("eoh_stress", EOH_STRESS))
        )

        st.divider()
        with st.expander("CAPEX S-Curve preview", expanded=False):
            a, b = st.columns(2)
            inflection = a.number_input(
                "Inflection Point",
                0.0,
                1.0,
                float(d("capex_curve_inflection", CAPEX_CURVE_INFLECTION)),
                0.05,
            )
            steepness = b.number_input(
                "Steepness",
                min_value=0.01,
                value=float(d("capex_curve_steepness", CAPEX_CURVE_STEEPNESS)),
            )
            monthly = np.asarray(
                distribute_capex(
                    nominal_capex,
                    int(nominal_time_to_cod),
                    inflection,
                    steepness,
                )
            )
            months = np.arange(1, len(monthly) + 1)
            fig = make_subplots(specs=[[{"secondary_y": True}]])
            fig.add_trace(go.Bar(x=months, y=monthly, name="Monthly CAPEX", marker=dict(color=CHART_BLUE_LIGHT, line=dict(color=BAR_OUTLINE, width=1.0)), hovertemplate="Month %{x}<br>Monthly CAPEX: %{y:.3f} M€<extra></extra>"), secondary_y=False)
            fig.add_trace(go.Scatter(x=months, y=np.cumsum(monthly), name="Cumulative CAPEX", mode="lines", line=dict(color=CHART_BLUE, width=3), hovertemplate="Month %{x}<br>Cumulative CAPEX: %{y:.3f} M€<extra></extra>"), secondary_y=True)
            style_figure(fig, "CAPEX S-Curve", "Month", "Monthly CAPEX (M€)", 470)
            fig.update_yaxes(title_text="Cumulative CAPEX (M€)", secondary_y=True)
            show_plotly(fig, "capex_s_curve")
    with operation:
        st.markdown("### Energy Price")
        energy_price = st.number_input(
            "Energy Price at COD (EUR/MWh)",
            min_value=0.0,
            value=float(d("energy_price_at_cod", ENERGY_PRICE_AT_COD)),
        )
        annual_price = st.number_input(
            "Annual Price Variation (EUR/MWh/y)",
            value=float(d("annual_price_variation", ANNUAL_PRICE_VARIATION)),
        )
        file_e = st.file_uploader(
            "Import Energy Price Excel", type=["xlsx"], key="e_upload"
        )
        if file_e is not None:
            st.session_state.energy_price_profile = parse_profile(file_e, "Energy Price")
        if st.button("Use Energy Price Model"):
            st.session_state.energy_price_profile = None
        e_values = energy_price_generator(
            0,
            int(operational_life),
            energy_price,
            annual_price,
            series=st.session_state.energy_price_profile,
        )
        st.download_button(
            "Download Energy Price Excel",
            profile_workbook("Energy Price", "EUR/MWh", e_values),
            "energy_price_profile.xlsx",
        )
        price_months = np.arange(1, len(e_values) + 1)
        fig = go.Figure(go.Scatter(x=price_months, y=e_values, name="Energy Price from Excel" if st.session_state.energy_price_profile is not None else "Energy Price Model", mode="lines", line=dict(color=CHART_BLUE, width=3), hovertemplate="Operating month %{x}<br>Price: %{y:.2f} EUR/MWh<extra></extra>"))
        style_figure(fig, "Monthly Energy Price Profile", "Operating Month", "Energy Price (EUR/MWh)")
        show_plotly(fig, "energy_price_profile")

        st.divider()
        st.markdown("### OPEX")
        opex = st.number_input(
            "OPEX (M€/MW/y)",
            min_value=0.0,
            value=float(d("opex", OPEX)),
            format="%.5f",
        )
        file_o = st.file_uploader("Import OPEX Excel", type=["xlsx"], key="o_upload")
        if file_o is not None:
            st.session_state.opex_profile = parse_profile(file_o, "OPEX")
        if st.button("Use OPEX Model"):
            st.session_state.opex_profile = None
        o_values = opex_generator(
            0,
            int(operational_life),
            opex,
            capacity,
            series=st.session_state.opex_profile,
        )
        st.download_button(
            "Download OPEX Excel",
            profile_workbook("OPEX", "M€/month", o_values),
            "opex_profile.xlsx",
        )
        opex_months = np.arange(1, len(o_values) + 1)
        fig = go.Figure(go.Scatter(x=opex_months, y=o_values, name="OPEX from Excel" if st.session_state.opex_profile is not None else "OPEX Model", mode="lines", # fill="tozeroy",
        line=dict(color=CHART_BLUE, width=3), #fillcolor="rgba(62,108,138,0.18)",
        hovertemplate="Operating month %{x}<br>OPEX: %{y:.4f} M€/month<extra></extra>"))
        style_figure(fig, "Monthly OPEX Profile", "Operating Month", "OPEX (M€/month)")
        show_plotly(fig, "opex_profile")

        st.divider()
        st.markdown("### Unavailability")
        u0 = st.number_input(
            "Initial Unavailability (%)",
            min_value=0.0,
            value=float(d("unavailability_parameters", UNAVAILABILITY_PARAMETERS)[0]),
        )
        us = st.number_input(
            "Steady-state Unavailability (%)",
            min_value=0.0,
            value=float(d("unavailability_parameters", UNAVAILABILITY_PARAMETERS)[1]),
        )
        ur = st.number_input(
            "Ramp-down (m)",
            min_value=1,
            value=int(d("unavailability_parameters", UNAVAILABILITY_PARAMETERS)[2]),
            step=1,
        )
        ua = st.number_input(
            "Ageing Degradation (%/m)",
            min_value=0.0,
            value=float(d("unavailability_parameters", UNAVAILABILITY_PARAMETERS)[3]),
            format="%.4f",
        )
        file_u = st.file_uploader(
            "Import Unavailability Excel", type=["xlsx"], key="u_upload"
        )
        if file_u is not None:
            st.session_state.unavailability_profile = parse_profile(
                file_u, "Unavailability"
            )
        if st.button("Use Unavailability Model"):
            st.session_state.unavailability_profile = None
        u_values = unavailability_generator(
            0,
            int(operational_life),
            u0,
            us,
            ur,
            ua,
            series=st.session_state.unavailability_profile,
        )
        st.download_button(
            "Download Unavailability Excel",
            profile_workbook("Unavailability", "%", np.asarray(u_values) * 100),
            "unavailability_profile.xlsx",
        )
        unavailability_months = np.arange(1, len(u_values) + 1)
        fig = go.Figure(go.Scatter(x=unavailability_months, y=np.asarray(u_values) * 100, name="Unavailability from Excel" if st.session_state.unavailability_profile is not None else "Unavailability Model", mode="lines", line=dict(color=CHART_BLUE, width=3), hovertemplate="Operating month %{x}<br>Unavailability: %{y:.2f}%<extra></extra>"))
        style_figure(fig, "Monthly Unavailability Profile", "Operating Month", "Unavailability (%)")
        show_plotly(fig, "unavailability_profile")
    with uncertainty:
        correlation = st.number_input(
            "CAPEX-Delay Correlation",
            min_value=-1.0,
            max_value=1.0,
            value=float(d("capex_delay_correlation", CAPEX_DELAY_CORRELATION)),
            step=0.05,
        )
        preview_iterations = st.number_input(
            "Preview Distribution Iterations",
            min_value=100,
            value=3000,
            step=100,
            key="preview_distribution_iterations",
        )
        st.divider()
        st.divider()
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("#### Delay")
            delay_min = st.number_input("Minimum Delay (m)", value=int(d("delay_min", DELAY_MIN)), step=1)
            delay_ml = st.number_input("Most Likely Delay (m)", value=float(d("delay_ml", DELAY_ML)))
            delay_max = st.number_input("Maximum Delay (m)", value=int(d("delay_max", DELAY_MAX)), step=1)
            delay_sharpness = st.number_input("Delay Distribution Sharpness", min_value=0.01, value=float(d("delay_distribution_sharpness", DELAY_DISTRIBUTION_SHARPNESS)))
            if stochastic_delay:
                values = np.asarray(time_to_cod_generator(int(nominal_time_to_cod), int(delay_min), int(delay_max), delay_ml, delay_sharpness, int(preview_iterations))) - nominal_time_to_cod
                fig = interactive_histogram(values, 0, "Delay Distribution Preview", "Delay (m)", discrete=True)
                show_plotly(fig, "delay_preview")
        with c2:
            st.markdown("#### CAPEX variation")
            capex_min = st.number_input("Minimum Variation (%)", value=float(d("capex_variation_min", CAPEX_VARIATION_MIN)))
            capex_ml = st.number_input("Most Likely Variation (%)", value=float(d("capex_variation_ml", CAPEX_VARIATION_ML)))
            capex_max = st.number_input("Maximum Variation (%)", value=float(d("capex_variation_max", CAPEX_VARIATION_MAX)))
            capex_sharpness = st.number_input("CAPEX Distribution Sharpness", min_value=0.01, value=float(d("capex_distribution_sharpness", CAPEX_DISTRIBUTION_SHARPNESS)))
            if stochastic_capex:
                values = (np.asarray(capex_generator(nominal_capex, capex_min, capex_max, capex_ml, capex_sharpness, int(preview_iterations))) / nominal_capex - 1) * 100
                fig = interactive_histogram(values, 0, "CAPEX Variation Distribution Preview", "CAPEX Variation (%)", 35)
                show_plotly(fig, "capex_preview")
        st.divider()
        st.markdown("### EOH")
        eoh_df = pd.DataFrame(
            {
                "Month": MONTHS,
                "Mean EOH": st.session_state.mean_eoh,
                "CV": st.session_state.std_eoh,
            }
        )
        edited = st.data_editor(
            eoh_df,
            hide_index=True,
            use_container_width=True,
            disabled=["Month"],
            key="eoh_editor",
        )
        st.session_state.mean_eoh = edited["Mean EOH"].astype(float).tolist()
        st.session_state.std_eoh = edited["CV"].astype(float).tolist()
        mean_eoh_variation = st.number_input(
            "Mean EOH Factor",
            min_value=0.0,
            value=float(d("mean_eoh_variation", MEAN_EOH_VARIATION)),
            format="%.4f",
        )
        means = np.asarray(st.session_state.mean_eoh)
        cvs = np.asarray(st.session_state.std_eoh)
        sigma = means * cvs
        fig = go.Figure()
        for mult, name, color in [(3, "±3σ", "rgba(138,182,204,0.25)"), (2, "±2σ", "rgba(119,101,154,0.25)"), (1, "±1σ", "rgba(63,131,122,0.30)")]:
            lower = np.maximum(0, means - mult * sigma)
            upper = means + mult * sigma
            fig.add_trace(go.Scatter(x=MONTHS, y=lower, mode="lines", line=dict(width=0), showlegend=False, hoverinfo="skip"))
            fig.add_trace(go.Scatter(x=MONTHS, y=upper, mode="lines", line=dict(width=0), fill="tonexty", fillcolor=color, name=name, hovertemplate=f"{name}: %{{y:.1f}} h<extra></extra>"))
        fig.add_trace(go.Scatter(x=MONTHS, y=means, name="Mean EOH", mode="lines+markers", line=dict(color=CHART_BLUE, width=3), marker=dict(size=7), hovertemplate="%{x}<br>Mean EOH: %{y:.1f} h<extra></extra>"))
        style_figure(fig, "Monthly EOH Profile", "Month", "EOH (h)")
        show_plotly(fig, "eoh_profile")


    with advanced:
        c1, c2, c3 = st.columns(3)
        irr_guess = c1.number_input("IRR Initial Guess", value=float(d("irr_guess", IRR_GUESS)), format="%.6f")
        irr_tol = c2.number_input("IRR Tolerance", value=float(d("irr_tol", IRR_TOL)), format="%.12f")
        irr_max_iter = c3.number_input("Maximum IRR Iterations", min_value=1, value=int(d("irr_max_iter", IRR_MAX_ITER)), step=1)

    return validate({
        "project_name": project_name, "iterations": int(iterations), "capex_delay_correlation": correlation,
        "stochastic_capex": stochastic_capex, "stochastic_delay": stochastic_delay, "stochastic_eoh": stochastic_eoh,
        "nominal_capex": nominal_capex, "start_year": int(start_year), "start_month": int(start_month),
        "nominal_time_to_cod": int(nominal_time_to_cod), "operational_life": int(operational_life), "capacity": capacity,
        "mean_eoh": st.session_state.mean_eoh, "std_eoh": st.session_state.std_eoh, "mean_eoh_variation": mean_eoh_variation,
        "unavailability_parameters": [u0, us, int(ur), ua], "opex": opex, "energy_price_at_cod": energy_price,
        "annual_price_variation": annual_price, "energy_price_profile": st.session_state.energy_price_profile,
        "unavailability_profile": st.session_state.unavailability_profile, "opex_profile": st.session_state.opex_profile,
        "tax_rate": tax_rate, "capex_depreciation_ratio": depreciation_ratio, "capex_depreciation_time": int(depreciation_time),
        "delay_min": int(delay_min), "delay_ml": delay_ml, "delay_max": int(delay_max), "delay_distribution_sharpness": delay_sharpness,
        "capex_variation_min": capex_min, "capex_variation_ml": capex_ml, "capex_variation_max": capex_max,
        "capex_distribution_sharpness": capex_sharpness, "capex_curve_inflection": inflection, "capex_curve_steepness": steepness,
        "capex_stress": capex_stress, "delay_stress": delay_stress, "eoh_stress": eoh_stress,
        "irr_guess": irr_guess, "irr_tol": irr_tol, "irr_max_iter": int(irr_max_iter),
    })


def run_monte_carlo(params):
    accepted = inspect.signature(my_irr_mc_function).parameters
    result = my_irr_mc_function(**{k: v for k, v in params.items() if k in accepted})
    return MonteCarloResults.from_tuple(result)


def run_single(params):
    accepted = inspect.signature(generate_flows).parameters
    return generate_flows(**{k: v for k, v in params.items() if k in accepted})


def summary(results):
    irr, pbp = finite(results.iterations_irr), finite(results.iterations_bpb)
    metrics = [
        ("Reference IRR", f"{results.reference_irr:.2%}"), ("Mean IRR", f"{np.mean(irr):.2%}"),
        ("Median IRR", f"{np.median(irr):.2%}"), ("P5 IRR", f"{np.percentile(irr,5):.2%}"), ("P95 IRR", f"{np.percentile(irr,95):.2%}"),
        ("Reference PBP", f"{results.reference_bpb:.2f} y"), ("Mean PBP", f"{np.mean(pbp):.2f} y"),
        ("Median PBP", f"{np.median(pbp):.2f} y"), ("P5 PBP", f"{np.percentile(pbp,5):.2f} y"), ("P95 PBP", f"{np.percentile(pbp,95):.2f} y"),
        ("Delay Stiffness", f"{results.delay_stiffness:.0f}"), ("CAPEX Stiffness", f"{results.capex_stiffness:.0f}"), ("EOH Stiffness", f"{results.eoh_stiffness:.0f}"),
    ]
    cols = st.columns(5)
    for i, (label, value) in enumerate(metrics): cols[i % 5].metric(label, value)


def results_tabs(results):
    t1, t2, t3, t4 = st.tabs(["IRR & PBP Histograms", "Monte Carlo vs Linear Model", "Other Histograms", "Scatter Plot"])
    with t1:
        bins = st.slider("Bins", 5, 150, 35, key="main_bins")
        show_plotly(interactive_histogram(np.asarray(results.iterations_irr) * 100, results.reference_irr * 100, "IRR Distribution", "IRR (%)", bins), "irr_hist")
        show_plotly(interactive_histogram(results.iterations_bpb, results.reference_bpb, "PBP Distribution", "PBP (y)", bins), "pbp_hist")
    with t2:
        x=np.asarray(results.percentiles_points); mc=np.asarray(results.irr_percentiles); linear=np.asarray(results.linear_irr_percentiles); error=np.asarray(results.linear_irr_error)
        fig=make_subplots(rows=2, cols=1, shared_xaxes=True, vertical_spacing=0.12, subplot_titles=("Monte Carlo vs Linear Model", "Linear Model Error"))
        fig.add_trace(go.Scatter(x=x,y=mc,name="Monte Carlo",line=dict(color=PRIMARY,width=3),hovertemplate="Percentile %{x}<br>MC: %{y:.2f} bps<extra></extra>"),row=1,col=1)
        fig.add_trace(go.Scatter(x=x,y=linear,name="Linear Model",line=dict(color=ACCENT,width=3,dash="dash"),hovertemplate="Percentile %{x}<br>Linear: %{y:.2f} bps<extra></extra>"),row=1,col=1)
        fig.add_trace(go.Scatter(x=x,y=error,name="Error",fill="tozeroy",line=dict(color=CHART_RED,width=2),fillcolor="rgba(180,76,76,.12)",hovertemplate="Percentile %{x}<br>Error: %{y:.2f} bps<extra></extra>"),row=2,col=1)
        style_figure(fig,"Percentile Comparison","Percentile","IRR Variation (bps)",700)
        fig.update_yaxes(title_text="Error (bps)",row=2,col=1)
        show_plotly(fig,"mc_linear")
    with t3:
        c1,c2=st.columns([3,1]); choice=c1.selectbox("Variable",list(DISTRIBUTIONS)); bins=c2.number_input("Bins",5,150,35)
        attr,xlabel,discrete=DISTRIBUTIONS[choice]; refs={"EOH":results.reference_eoh,"EOH Variation (%)":0,"Time to COD (m)":results.reference_time_to_cod,"Delay (m)":0,"CAPEX (M€)":results.reference_capex,"CAPEX Variation (%)":0,"IRR Variation (bps)":0}
        show_plotly(interactive_histogram(getattr(results,attr),refs[choice],choice,xlabel,int(bins),discrete),"other_hist")
    with t4:
        c1,c2=st.columns(2); xname=c1.selectbox("X axis",list(SCATTER_VARIABLES),index=5); yname=c2.selectbox("Y axis",list(SCATTER_VARIABLES),index=7)
        xa,xl,xs=SCATTER_VARIABLES[xname]; ya,yl,ys=SCATTER_VARIABLES[yname]
        x=np.asarray(getattr(results,xa),float)*xs; y=np.asarray(getattr(results,ya),float)*ys; valid=np.isfinite(x)&np.isfinite(y); x,y=x[valid],y[valid]
        fig=go.Figure(go.Scattergl(x=x,y=y,mode="markers",name="Simulations",marker=dict(color=CHART_BLUE,size=6,opacity=.38,line=dict(color=BAR_OUTLINE,width=.35)),hovertemplate=f"{xl}: %{{x:.3f}}<br>{yl}: %{{y:.3f}}<extra></extra>"))
        if len(x)>1 and np.ptp(x)>0:
            slope,intercept,r,_,_=linregress(x,y); xx=np.linspace(x.min(),x.max(),200)
            fig.add_trace(go.Scatter(x=xx,y=intercept+slope*xx,name=f"Linear fit (r={r:.3f})",line=dict(color=ACCENT,width=3)))
        style_figure(fig,f"{yname} vs {xname}",xl,yl,560,"closest")
        show_plotly(fig,"scatter")

def single_iteration_view(generated, params):
    irr, actual_eoh, pbp, capex, eoh, unavailability, price, revenues, taxes, opex = generated
    arrays = [np.asarray(v,float) for v in (capex,eoh,unavailability,price,revenues,taxes,opex)]
    capex,eoh,unavailability,price,revenues,taxes,opex = arrays
    total=len(capex); cod=total-int(params["operational_life"])*12; months=np.arange(1,total+1); net=-capex+revenues-opex-taxes; cumulative=np.cumsum(net)
    c=st.columns(5); c[0].metric("IRR",f"{irr:.2%}"); c[1].metric("PBP",f"{pbp:.2f} y"); c[2].metric("Actual CAPEX",f"{capex.sum():.2f} M€"); c[3].metric("Actual Time to COD",f"{cod} m"); c[4].metric("Total EOH",f"{actual_eoh:,.0f} h")
    st.download_button(
        "Download Single Iteration Excel",
        single_iteration_workbook(generated, params),
        "single_iteration_results.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
        key="download_single_iteration_excel",
    )
    chart, table = st.tabs(["Charts","Monthly Flows"])
    with chart:
        # 1. Financial flows
        fig = make_subplots(specs=[[{"secondary_y": True}]])
        fig.add_trace(
            go.Bar(
                x=months,
                y=net,
                name="Net Cash Flow",
                marker=dict(color=CHART_BLUE, line=dict(color=BAR_OUTLINE, width=1.0)),
                opacity=0.82,
                hovertemplate="Month %{x}<br>Net cash flow: %{y:.3f} M€<extra></extra>",
            ),
            secondary_y=False,
        )
        fig.add_trace(
            go.Scatter(
                x=months,
                y=cumulative,
                name="Cumulative Cash Flow",
                mode="lines",
                line=dict(color=CHART_ORANGE, width=3),
                hovertemplate="Month %{x}<br>Cumulative cash flow: %{y:.3f} M€<extra></extra>",
            ),
            secondary_y=True,
        )
        fig.add_hline(y=0, line_color=MUTED, line_width=1)
        fig.add_vline(x=cod, line_dash="dash", line_color=TEXT)
        style_figure(fig, "Financial Flows", "Month", "Net Cash Flow (M€)", 500)
        fig.update_yaxes(title_text="Cumulative Cash Flow (M€)", secondary_y=True)
        show_plotly(fig, "single_financial_flows")

        # 2. EOH and unavailability
        fig = make_subplots(specs=[[{"secondary_y": True}]])
        fig.add_trace(
            go.Scatter(
                x=months,
                y=eoh,
                name="EOH",
                mode="lines",
                line=dict(color=CHART_BLUE, width=2.5),
                hovertemplate="Month %{x}<br>EOH: %{y:.2f} h<extra></extra>",
            ),
            secondary_y=False,
        )
        fig.add_trace(
            go.Scatter(
                x=months,
                y=unavailability * 100,
                name="Unavailability",
                mode="lines",
                line=dict(color=CHART_ORANGE, width=2.5),
                hovertemplate="Month %{x}<br>Unavailability: %{y:.2f}%<extra></extra>",
            ),
            secondary_y=True,
        )
        fig.add_vline(x=cod, line_dash="dash", line_color=TEXT)
        style_figure(fig, "EOH and Unavailability", "Month", "EOH (h)", 500)
        fig.update_yaxes(title_text="Unavailability (%)", secondary_y=True)
        show_plotly(fig, "single_eoh_unavailability")

        # 3. Energy price and revenues
        fig = make_subplots(specs=[[{"secondary_y": True}]])
        fig.add_trace(
            go.Scatter(
                x=months,
                y=price,
                name="Energy Price",
                mode="lines",
                line=dict(color=CHART_BLUE, width=2.5),
                hovertemplate="Month %{x}<br>Energy price: %{y:.2f} EUR/MWh<extra></extra>",
            ),
            secondary_y=False,
        )
        fig.add_trace(
            go.Scatter(
                x=months,
                y=revenues,
                name="Gross Revenues",
                mode="lines",
                line=dict(color=CHART_ORANGE, width=2.5),
                hovertemplate="Month %{x}<br>Gross revenues: %{y:.3f} M€<extra></extra>",
            ),
            secondary_y=True,
        )
        fig.add_vline(x=cod, line_dash="dash", line_color=TEXT)
        style_figure(fig, "Energy Price and Revenues", "Month", "Energy Price (EUR/MWh)", 500)
        fig.update_yaxes(title_text="Gross Revenues (M€)", secondary_y=True)
        show_plotly(fig, "single_price_revenues")

        # 4. Overlaid revenue and cost bridge
        opex_base = revenues
        capex_base = revenues - opex
        bar_width = 0.76

        fig = go.Figure()
        fig.add_trace(
            go.Bar(
                x=months,
                y=revenues,
                base=0,
                width=bar_width,
                name="Gross Revenues",
                marker=dict(color="#8FC5A0", line=dict(color=BAR_OUTLINE, width=1.1)),
                opacity=0.82,
                hovertemplate="Month %{x}<br>Gross revenues: %{y:.3f} M€<extra></extra>",
            )
        )
        fig.add_trace(
            go.Bar(
                x=months,
                y=taxes,
                base=0,
                width=bar_width * 0.56,
                name="Taxes",
                marker=dict(color=CHART_GOLD, line=dict(color=BAR_OUTLINE, width=1.0)),
                opacity=0.92,
                hovertemplate="Month %{x}<br>Taxes: %{y:.3f} M€<extra></extra>",
            )
        )
        fig.add_trace(
            go.Bar(
                x=months,
                y=-opex,
                base=opex_base,
                width=bar_width * 0.62,
                name="OPEX",
                marker=dict(color=CHART_ORANGE, line=dict(color=BAR_OUTLINE, width=1.0)),
                opacity=0.95,
                customdata=opex,
                hovertemplate="Month %{x}<br>OPEX: %{customdata:.3f} M€<extra></extra>",
            )
        )
        fig.add_trace(
            go.Bar(
                x=months,
                y=-capex,
                base=capex_base,
                width=bar_width * 0.48,
                name="CAPEX",
                marker=dict(color=CHART_SLATE, line=dict(color=BAR_OUTLINE, width=1.0)),
                opacity=0.96,
                customdata=capex,
                hovertemplate="Month %{x}<br>CAPEX: %{customdata:.3f} M€<extra></extra>",
            )
        )
        fig.add_hline(y=0, line_color=MUTED, line_width=1)
        fig.add_vline(x=cod, line_dash="dash", line_color=TEXT)
        style_figure(
            fig,
            "Revenues, Taxes, OPEX and CAPEX",
            "Month",
            "Monthly Flow (M€)",
            560,
            "x unified",
        )
        fig.update_layout(barmode="overlay", bargap=0.08)
        show_plotly(fig, "single_revenue_cost_bridge")
    with table:
        dates=[date(int(params["start_year"])+(int(params["start_month"])-1+i)//12,(int(params["start_month"])-1+i)%12+1,15) for i in range(total)]
        df=pd.DataFrame({"Month":months,"Date":dates,"Phase":["Construction" if i<cod else "Operation" for i in range(total)],"Operating Month":[0 if i<cod else i-cod+1 for i in range(total)],"CAPEX (M€)":capex,"EOH (h)":eoh,"Unavailability":unavailability,"Energy Price":price,"Revenues (M€)":revenues,"OPEX (M€)":opex,"Taxes (M€)":taxes,"Net CF (M€)":net,"Cumulative CF (M€)":cumulative})
        st.dataframe(df,use_container_width=True,height=550)


def apply_supabase_preset(row):
    """Apply one Supabase preset without assigning any file-uploader key."""
    parameters = row.get("parameters", {})
    if not isinstance(parameters, dict):
        raise ValueError("The preset does not contain a valid parameters object.")

    parameters = dict(parameters)
    mean_eoh = list(parameters.get("mean_eoh", MEAN_EOH))
    std_eoh = list(parameters.get("std_eoh", STD_EOH))
    if len(mean_eoh) != 12 or len(std_eoh) != 12:
        raise ValueError("Mean EOH and EOH CV must contain exactly 12 values.")

    # Remove every prior widget value, including complete_inputs_upload.
    # The uploader key is deliberately NOT restored or assigned: Streamlit owns it.
    for key in list(st.session_state.keys()):
        del st.session_state[key]

    st.session_state.imported_inputs = parameters
    st.session_state.mean_eoh = [float(value) for value in mean_eoh]
    st.session_state.std_eoh = [float(value) for value in std_eoh]
    st.session_state.energy_price_profile = (
        list(parameters["energy_price_profile"])
        if parameters.get("energy_price_profile") is not None
        else None
    )
    st.session_state.opex_profile = (
        list(parameters["opex_profile"])
        if parameters.get("opex_profile") is not None
        else None
    )
    st.session_state.unavailability_profile = (
        list(parameters["unavailability_profile"])
        if parameters.get("unavailability_profile") is not None
        else None
    )
    st.session_state.selected_preset_id = int(row["id"])
    st.session_state.selected_preset_name = str(row.get("name", ""))
    st.session_state.selected_preset_description = str(
        row.get("description", "") or ""
    )
    st.session_state.status = "Inputs loaded from Supabase preset"
    st.session_state.error = None


def supabase_preset_manager(params):
    """Render a single preset manager after the original Excel/input section."""
    st.markdown("### Supabase presets")
    try:
        test_connection()
        presets = list_presets()
    except Exception as exc:
        st.warning(f"Supabase presets unavailable: {exc}")
        return

    selected_id = None
    if presets:
        rows_by_id = {int(row["id"]): row for row in presets}
        preset_ids = list(rows_by_id)
        current_id = st.session_state.get("selected_preset_id")
        default_index = preset_ids.index(current_id) if current_id in preset_ids else 0
        selected_id = st.selectbox(
            "Saved preset",
            preset_ids,
            index=default_index,
            format_func=lambda value: (
                f"{rows_by_id[value].get('name', 'Unnamed')} (ID {value})"
            ),
            key="db_preset_selector",
        )

        load_column, delete_column = st.columns(2)
        if load_column.button(
            "Load preset", key="db_load_preset", use_container_width=True
        ):
            try:
                apply_supabase_preset(get_preset(selected_id))
                st.rerun()
            except Exception as exc:
                st.error(f"Unable to load preset: {exc}")

        if delete_column.button(
            "Delete preset", key="db_delete_preset", use_container_width=True
        ):
            try:
                delete_preset(selected_id)
                st.session_state.selected_preset_id = None
                st.session_state.selected_preset_name = ""
                st.session_state.selected_preset_description = ""
                st.session_state.status = "Preset deleted"
                st.rerun()
            except Exception as exc:
                st.error(f"Unable to delete preset: {exc}")
    else:
        st.info("No Supabase presets saved yet.")

    preset_name = st.text_input(
        "Preset name",
        value=st.session_state.get("selected_preset_name", ""),
        key="db_preset_name",
    )
    preset_description = st.text_area(
        "Preset description",
        value=st.session_state.get("selected_preset_description", ""),
        key="db_preset_description",
    )

    save_column, update_column = st.columns(2)
    if save_column.button(
        "Save as new preset",
        type="primary",
        key="db_save_preset",
        use_container_width=True,
    ):
        try:
            row = save_preset(preset_name, params, preset_description)
            st.session_state.selected_preset_id = int(row["id"])
            st.session_state.selected_preset_name = str(
                row.get("name", preset_name)
            )
            st.session_state.selected_preset_description = str(
                row.get("description", preset_description) or ""
            )
            st.session_state.status = "Preset saved"
            st.rerun()
        except Exception as exc:
            st.error(f"Unable to save preset: {exc}")

    if update_column.button(
        "Update selected preset",
        key="db_update_preset",
        disabled=st.session_state.get("selected_preset_id") is None,
        use_container_width=True,
    ):
        try:
            row = update_preset(
                st.session_state.selected_preset_id,
                preset_name,
                params,
                preset_description,
            )
            st.session_state.selected_preset_name = str(
                row.get("name", preset_name)
            )
            st.session_state.selected_preset_description = str(
                row.get("description", preset_description) or ""
            )
            st.session_state.status = "Preset updated"
            st.rerun()
        except Exception as exc:
            st.error(f"Unable to update preset: {exc}")


def main():
    init_state()
    status_class = "status-error" if st.session_state.error else "status-ready"
    st.markdown(
        f'<div class="hero"><h1>IRR Monte Carlo DB</h1>'
        f'<p>Project-finance simulation and risk-analysis tool</p></div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        f'<span class="{status_class}">{st.session_state.status}</span>',
        unsafe_allow_html=True,
    )

    input_tab, monte_carlo_tab, single_iteration_tab = st.tabs(
        ["Inputs", "Monte Carlo", "Single Iteration"]
    )

    with input_tab:
        st.markdown("### Load complete input workbook")
        upload_column, template_column = st.columns([3, 2])
        with upload_column:
            uploaded_inputs = st.file_uploader(
                "Import Inputs and Operating Profiles",
                type=["xlsx"],
                key="complete_inputs_upload",
            )
        with template_column:
            st.markdown("#### Input workbook template")
            template_params = st.session_state.get("current_params")
            if template_params is None:
                template_params = input_defaults()
            st.download_button(
                "Download Input Template",
                params_workbook(template_params),
                "irr_inputs_template.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="download_input_template",
            )
        if st.button(
            "Load Inputs from Excel",
            disabled=uploaded_inputs is None,
            key="load_complete_inputs",
        ):
            try:
                imported_inputs, imported_profiles, imported_flags = (
                    load_params_workbook(uploaded_inputs)
                )
                # Do not assign to or recreate the file_uploader widget key.
                # Streamlit owns complete_inputs_upload and forbids setting it
                # through st.session_state.
                keys_to_clear = [
                    key
                    for key in st.session_state.keys()
                    if key not in {"complete_inputs_upload"}
                ]
                for key in keys_to_clear:
                    del st.session_state[key]
                st.session_state.imported_inputs = imported_inputs
                st.session_state.energy_price_profile = (
                    list(imported_profiles["energy_price"])
                    if imported_flags["energy_price"]
                    else None
                )
                st.session_state.opex_profile = (
                    list(imported_profiles["opex"])
                    if imported_flags["opex"]
                    else None
                )
                st.session_state.unavailability_profile = (
                    list(imported_profiles["unavailability"])
                    if imported_flags["unavailability"]
                    else None
                )
                st.session_state.mean_eoh = list(imported_inputs["mean_eoh"])
                st.session_state.std_eoh = list(imported_inputs["std_eoh"])
                st.session_state.status = "Inputs loaded from Excel"
                st.session_state.error = None
                st.rerun()
            except Exception as exc:
                st.error(f"Unable to load workbook: {exc}")

        st.divider()
        try:
            params = input_panel(st.session_state.get("imported_inputs"))
            st.session_state.current_params = dict(params)
            st.session_state.error = None
        except Exception as exc:
            st.session_state.error = f"{type(exc).__name__}: {exc}"
            st.session_state.status = "Invalid input"
            st.error(st.session_state.error)
            return


        st.divider()
        supabase_preset_manager(params)
    params = st.session_state.get("current_params")

    with monte_carlo_tab:
        st.markdown("## Monte Carlo Simulation")
        if params is None:
            st.info("Complete the Inputs tab before running the simulation.")
        else:
            if st.button(
                "Run Monte Carlo",
                type="primary",
                use_container_width=True,
                key="run_monte_carlo",
            ):
                try:
                    with st.spinner("Simulation running..."):
                        st.session_state.results = run_monte_carlo(params)
                        st.session_state.monte_carlo_params = dict(params)
                    st.session_state.status = (
                        f"Completed: "
                        f"{len(st.session_state.results.iterations_irr):,} iterations"
                    )
                    st.session_state.error = None
                except Exception as exc:
                    st.session_state.error = f"{type(exc).__name__}: {exc}"
                    st.session_state.status = "Simulation error"

            if st.session_state.error:
                st.error(st.session_state.error)

            if st.session_state.results is not None:
                st.markdown("## Simulation Summary")
                summary(st.session_state.results)
                results_tabs(st.session_state.results)
            else:
                st.info("Run the Monte Carlo simulation to display its results.")

    with single_iteration_tab:
        st.markdown("## Single Iteration")
        if params is None:
            st.info("Complete the Inputs tab before running a single iteration.")
        else:
            if st.button(
                "Run Single Iteration",
                type="primary",
                use_container_width=True,
                key="run_single_iteration",
            ):
                try:
                    with st.spinner("Single iteration running..."):
                        st.session_state.single_iteration = run_single(params)
                        st.session_state.single_iteration_params = dict(params)
                    st.session_state.status = "Single iteration completed"
                    st.session_state.error = None
                except Exception as exc:
                    st.session_state.error = f"{type(exc).__name__}: {exc}"
                    st.session_state.status = "Single iteration error"

            if st.session_state.error:
                st.error(st.session_state.error)

            if st.session_state.single_iteration is not None:
                single_iteration_view(
                    st.session_state.single_iteration,
                    st.session_state.single_iteration_params,
                )
            else:
                st.info("Run a single iteration to display its results.")


if __name__ == "__main__":
    main()
